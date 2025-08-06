import streamlit as st
from openpyxl import load_workbook
import yagmail
import tempfile
import os
from datetime import datetime
import pandas as pd
import traceback

def show_report_form(user_role="guest"):
    st.title("📋 Submit Project Report")

    entered_by = st.text_input("👤 Your Name").strip().title()
    if not entered_by:
        st.warning("Please enter your name to proceed.")
        return

    firm_name = st.text_input("🏢 Firm Name").upper()
    firm_address = st.text_input("📍 Firm Address").upper()
    nature_of_business = st.text_input("💼 Nature of Business").upper()

    building_status = st.selectbox("🏗️ Building Status", ["RENTED", "OWNED"])
    area_sqft = st.number_input("📐 Area in sq. ft", min_value=0.0, step=1.0)
    rent_rs = st.number_input("💰 Rent in Rs.", min_value=0.0, step=100.0) if building_status == "RENTED" else 0.0

    st.markdown("### 💸 Financial Parameters")
    margin = st.number_input("Margin %", min_value=0.0, max_value=100.0) / 100
    subsidy = st.number_input("Subsidy %", min_value=0.0, max_value=100.0) / 100
    term_loan = st.number_input("Term Loan Interest %", min_value=0.0, max_value=100.0) / 100
    cc_rate = st.number_input("CC Interest %", min_value=0.0, max_value=100.0) / 100

    if st.button("📤 Submit & Generate Report"):
        try:
            template_path = "Project Report Format.xlsx"
            wb = load_workbook(template_path)

            if "Basic Details" not in wb.sheetnames:
                st.error("❌ 'Basic Details' sheet missing in the template.")
                return

            sheet = wb["Basic Details"]
            sheet["C3"] = firm_name
            sheet["C4"] = firm_address
            sheet["C5"] = nature_of_business
            sheet["C8"] = building_status
            sheet["C9"] = area_sqft
            sheet["C10"] = rent_rs
            sheet["C11"] = margin
            sheet["C12"] = subsidy
            sheet["C13"] = term_loan
            sheet["C14"] = cc_rate

            # 🔒 Hide sheet
            sheet.sheet_state = "hidden"

            # Save to temp file
            now = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_name = entered_by.replace(" ", "_")
            filename = f"/tmp/report_{safe_name}_{now}.xlsx"
            wb.save(filename)

            # ✅ Send Email
            yag = yagmail.SMTP(user="calavishgupta25@gmail.com", password="geli tejz dxiq vtyo")
            subject = f"{entered_by} - Project Report"
            yag.send(
                to="calavishgupta25@gmail.com",
                subject=subject,
                contents="Attached is the project report submitted via the Streamlit app.",
                attachments=filename
            )
            st.success("✅ Report generated and emailed successfully.")

            # 📁 Log Submission
            new_row = pd.DataFrame([{
                "submitted_by": entered_by,
                "email": st.session_state.email if user_role != "guest" else "guest",
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "file_path": filename
            }])

            if os.path.exists("submissions.csv"):
                existing = pd.read_csv("submissions.csv")
                combined = pd.concat([existing, new_row], ignore_index=True)
            else:
                combined = new_row

            combined.to_csv("submissions.csv", index=False)

        except Exception as e:
            st.error(f"❌ Failed to generate or send report.")
            st.exception(e)  # Shows traceback for easier debugging

    # --- Debug Info ---
    with st.expander("🛠️ Debug Info"):
        st.write("Entered By:", entered_by)
        st.write("Firm Name:", firm_name)
        st.write("Nature of Business:", nature_of_business)
        st.write("Margin:", margin, "| Subsidy:", subsidy, "| TL Interest:", term_loan, "| CC Rate:", cc_rate)
